#!/usr/bin/env python3
"""Fill the bundled PChK PPTX. Never modify the template or silently truncate text."""
from __future__ import annotations

import argparse
import copy
import json
import math
import os
from pathlib import Path
import re
import shutil
import subprocess
import sys
import tempfile

from PIL import Image, ImageFont
from pptx import Presentation
from pptx.util import Pt

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "assets" / "character-sheet.pptx"
STAT_NAMES = dict(str="Сила", dex="Ловкость", con="Выносливость", wis="Мудрость", int="Интеллект", cha="Харизма")
BODY_SIZE = 9.4
LEADING = BODY_SIZE * 1.04


class SheetError(ValueError):
    pass


def write_json(path, value):
    with Path(path).open("w", encoding="utf-8", newline="\r\n") as stream:
        json.dump(value, stream, ensure_ascii=False, indent=2)
        stream.write("\n")


def obj(value, field, allowed, required=()):
    if not isinstance(value, dict):
        raise SheetError(f"{field}: ожидается объект")
    unknown, missing = set(value) - set(allowed), set(required) - set(value)
    if unknown or missing:
        raise SheetError(f"{field}: неизвестные поля {sorted(unknown)}; отсутствуют {sorted(missing)}")


def string(value, field, required=False):
    if not isinstance(value, str) or (required and not value.strip()):
        raise SheetError(f"{field}: ожидается {'непустая ' if required else ''}строка")
    if any(ord(c) < 32 and c not in "\n\r\t" for c in value):
        raise SheetError(f"{field}: недопустимые управляющие символы")


def integer(value, field, minimum, maximum):
    if type(value) is not int or not minimum <= value <= maximum:
        raise SheetError(f"{field}: нужно целое число от {minimum} до {maximum}")


def entries(value, field):
    if not isinstance(value, list):
        raise SheetError(f"{field}: ожидается список")
    for i, row in enumerate(value):
        key = f"{field}[{i}]"
        obj(row, key, ["name", "uses", "text"], ["name"])
        string(row["name"], key + ".name", True)
        for k in ["uses", "text"]:
            if k in row:
                string(row[k], key + "." + k)


def load_draft(path):
    path = Path(path).resolve()
    try:
        data = json.loads(path.read_text(encoding="utf-8-sig"))
    except (OSError, ValueError) as exc:
        raise SheetError(f"Не удалось прочитать черновик: {exc}") from exc
    required = ["version", "name", "ability_modifiers", "armor_class", "hit_points", "abilities", "equipment", "backstory"]
    allowed = required + ["role", "rules", "weapon_proficiencies", "skills", "special", "portrait", "portrait_focus"]
    obj(data, "draft", allowed, required)
    integer(data["version"], "version", 1, 1)
    for k in ["name", "backstory"]:
        string(data[k], k, True)
    for k in ["role", "rules", "weapon_proficiencies", "skills"]:
        string(data.get(k, ""), k)
    obj(data["ability_modifiers"], "ability_modifiers", STAT_NAMES, STAT_NAMES)
    for k, value in data["ability_modifiers"].items():
        integer(value, "ability_modifiers." + k, -5, 15)
    ac = data["armor_class"]
    if type(ac) is int:
        integer(ac, "armor_class", 1, 99)
    elif not isinstance(ac, str) or not re.fullmatch(r"[1-9]\d?/[1-9]\d?", ac):
        raise SheetError("armor_class: нужно целое 1…99 либо строка вида 14/16")
    integer(data["hit_points"], "hit_points", 1, 9999)
    entries(data["abilities"], "abilities")
    entries(data["equipment"], "equipment")
    special = data.get("special")
    if special is not None:
        obj(special, "special", ["title", "charges", "recovery", "effects"], ["title", "effects"])
        string(special["title"], "special.title", True)
        if "charges" in special:
            integer(special["charges"], "special.charges", 1, 999)
        string(special.get("recovery", ""), "special.recovery")
        entries(special["effects"], "special.effects")
    focus = data.get("portrait_focus", [0.5, 0.5])
    if (not isinstance(focus, list) or len(focus) != 2
            or any(type(v) not in (int, float) or not math.isfinite(v) or not 0 <= v <= 1 for v in focus)):
        raise SheetError("portrait_focus: нужны два числа от 0 до 1")
    if data.get("portrait") is not None:
        string(data["portrait"], "portrait", True)
        portrait = Path(data["portrait"])
        portrait = portrait if portrait.is_absolute() else path.parent / portrait
        if not portrait.is_file():
            raise SheetError(f"Портрет не найден: {portrait}")
        try:
            with Image.open(portrait) as image:
                if image.format not in ("PNG", "JPEG", "WEBP"):
                    raise SheetError("Портрет должен быть PNG, JPEG или WebP")
                image.verify()
        except (OSError, ValueError) as exc:
            raise SheetError(f"Повреждённый портрет: {exc}") from exc
    return data


def rows(data):
    result = []
    for row in data:
        label = row["name"]
        if row.get("uses"):
            label += f" ({row['uses']})"
        desc = row.get("text", "")
        result.append((label + (": " if desc else ""), desc))
    return result


def shape_by_name(slide, name):
    found = [s for s in slide.shapes if s.name == name]
    if len(found) != 1:
        raise SheetError(f"Шаблон несовместим: нужен один элемент {name!r}, найдено {len(found)}")
    return found[0]


def remove(shape):
    shape._element.getparent().remove(shape._element)


def set_text(shape, text):
    """Preserve the first run's font/effects, including the outlined shield numbers."""
    p = shape.text_frame.paragraphs[0]
    proto = copy.deepcopy(p.runs[0]._r.rPr) if p.runs else None
    pproto = copy.deepcopy(p._p.pPr)
    shape.text_frame.clear()
    for i, line in enumerate(text.split("\n")):
        paragraph = shape.text_frame.paragraphs[0] if i == 0 else shape.text_frame.add_paragraph()
        if pproto is not None:
            if paragraph._p.pPr is not None:
                paragraph._p.remove(paragraph._p.pPr)
            paragraph._p.insert(0, copy.deepcopy(pproto))
        run = paragraph.add_run()
        run.text = line
        if proto is not None:
            run._r.insert(0, copy.deepcopy(proto))


def set_rows(shape, content):
    first = shape.text_frame.paragraphs[0]
    prototypes = [copy.deepcopy(run._r.rPr) for run in first.runs[:2]]
    if len(prototypes) != 2:
        raise SheetError(f"В шаблоне потеряны стили блока {shape.name}")
    shape.text_frame.clear()
    for i, (label, desc) in enumerate(content):
        p = shape.text_frame.paragraphs[0] if i == 0 else shape.text_frame.add_paragraph()
        p.space_before = p.space_after = Pt(0)
        p.line_spacing = Pt(LEADING)
        for j, value in enumerate((label, desc)):
            if value:
                run = p.add_run()
                run.text = value
                run._r.insert(0, copy.deepcopy(prototypes[j]))
                run.font.size = Pt(BODY_SIZE)


def set_number(shape, value):
    set_text(shape, value)
    shape.text_frame.word_wrap = False
    run = shape.text_frame.paragraphs[0].runs[0]
    size = run.font.size.pt
    font_path = Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts" / "impact.ttf"
    if font_path.is_file():
        font = ImageFont.truetype(str(font_path), round(size * 10))
        measured = font.getlength(value) / 10
        size = min(size, size * (shape.width.pt - 3) / max(measured, 1))
    elif len(value) > 2:
        size = min(size, 42 / (len(value) * 0.55))
    run.font.size = Pt(size)


def font_for(bold=False):
    directory = Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts"
    path = directory / ("cambriab.ttf" if bold else "cambriai.ttf")
    if path.is_file():
        return ImageFont.truetype(str(path), round(BODY_SIZE * 10))
    fallback = ROOT / "assets" / "fonts" / "ptserif" / ("PT_Serif-Web-Bold.ttf" if bold else "PT_Serif-Web-Italic.ttf")
    return ImageFont.truetype(str(fallback), round(BODY_SIZE * 10)) if fallback.is_file() else None


def text_height(content, width):
    """Approximate mixed-style wrapping; native PowerPoint checks the final result."""
    fonts = [font_for(True), font_for(False)]
    line_count = 0
    for label, desc in content:
        used, count = 0.0, 1
        for index, part in enumerate((label, desc)):
            for token in re.findall(r"\n|[^\S\n]+|\S+", part):
                if token == "\n":
                    count += 1
                    used = 0
                    continue
                size = fonts[index].getlength(token) / 10 if fonts[index] else len(token) * BODY_SIZE * 0.55
                if used + size > width and used > 0 and not token.isspace():
                    count += 1
                    used = 0
                if size > width:
                    count += max(0, math.ceil(size / width) - 1)
                    used = size % width
                else:
                    used += size
        line_count += count
    return max(LEADING, line_count * LEADING) + 2


def fill(draft_path, data):
    presentation = Presentation(str(TEMPLATE))
    if len(presentation.slides) != 2:
        raise SheetError("Ожидаются два слайда в шаблоне")
    stats, bio = presentation.slides
    title = ", ".join(filter(None, [data.get("role"), data["name"]]))
    set_text(shape_by_name(stats, "Имя и роль персонажа"), title)
    for key, name in STAT_NAMES.items():
        set_number(shape_by_name(stats, name), f"{data['ability_modifiers'][key]:+d}")
    set_number(shape_by_name(stats, "Класс брони"), str(data["armor_class"]))
    set_number(shape_by_name(stats, "Здоровье"), str(data["hit_points"]))

    summary = []
    for key, label in [("weapon_proficiencies", "Владение оружием: "), ("skills", "Особые навыки: ")]:
        if data.get(key):
            summary.append((label, data[key]))
    summary_shape = shape_by_name(stats, "Владения и особые навыки")
    y = 124.0
    if summary:
        set_rows(summary_shape, summary)
        hh = text_height(summary, 535)
        summary_shape.top, summary_shape.height = Pt(y), Pt(hh)
        y += hh + 9
    else:
        remove(summary_shape)

    special = data.get("special")
    special_title = ""
    effects = []
    if special:
        special_title = special["title"]
        if "charges" in special:
            special_title += f" (заряды: {special['charges']})"
        effects = rows(special["effects"])
        if special.get("recovery"):
            effects.append(("Восстановление: ", special["recovery"]))
    sections = [
        ("Заголовок способностей", "Способности персонажа", "Способности:", rows(data["abilities"])),
        ("Заголовок вещей", "Оружие, броня и инвентарь", "Вещи:", rows(data["equipment"])),
        ("Заголовок особого ресурса", "Особые способности и заряды", special_title + ":", effects),
    ]
    for heading_name, body_name, heading, content in sections:
        header, body = shape_by_name(stats, heading_name), shape_by_name(stats, body_name)
        if not content:
            remove(header)
            remove(body)
            continue
        set_text(header, heading)
        header.top, header.height = Pt(y), Pt(15)
        y += 15
        set_rows(body, content)
        hh = text_height(content, 535)
        body.top, body.height = Pt(y), Pt(hh)
        y += hh + 9
    if y - 9 > 407:
        raise SheetError(f"Первый лист переполнен примерно на {y - 9 - 407:.1f} пт. Сократите описания без потери механики.")

    set_text(shape_by_name(bio, "Предыстория персонажа"), data["backstory"])
    for shape in list(bio.shapes):
        if shape.name.startswith("Портрет"):
            remove(shape)
    if data.get("portrait"):
        path = Path(data["portrait"])
        path = path if path.is_absolute() else Path(draft_path).resolve().parent / path
        width, height = presentation.slide_width / 2, presentation.slide_height
        # Pillow only decodes formats here. Cropping is native PPTX metadata, not a raster edit.
        with Image.open(path) as image:
            iw, ih = image.size
            normalized = None
            if image.format == "WEBP":
                import io
                normalized = io.BytesIO()
                image.save(normalized, format="PNG")
                normalized.seek(0)
        pic = bio.shapes.add_picture(normalized or str(path), 0, 0, width=int(width), height=int(height))
        pic.name = "Портрет персонажа — заменяемое изображение"
        source_ratio, target_ratio = iw / ih, width / height
        fx, fy = data.get("portrait_focus", [0.5, 0.5])
        if source_ratio > target_ratio:
            fraction = target_ratio / source_ratio
            pic.crop_left = min(1 - fraction, max(0, fx - fraction / 2))
            pic.crop_right = 1 - fraction - pic.crop_left
        else:
            fraction = source_ratio / target_ratio
            pic.crop_top = min(1 - fraction, max(0, fy - fraction / 2))
            pic.crop_bottom = 1 - fraction - pic.crop_top
    notes = f"Персонаж: {data['name']}\nПравила и допущения: {data.get('rules', 'не указаны')}\nИсходный черновик: {Path(draft_path).name}"
    for slide in presentation.slides:
        slide.notes_slide.notes_text_frame.text += "\n\n" + notes
    presentation.core_properties.title = data["name"] + " — лист ПЧК"
    return presentation


def output_paths(pptx):
    return pptx.with_suffix(".pdf"), pptx.with_suffix(".preview"), pptx.with_suffix(".render-report.json")


def preflight(paths, force):
    for path in paths:
        if path.exists() and not force:
            raise SheetError(f"Результат уже существует: {path}. Для пересборки добавьте --force.")


def render(pptx, pdf=False, preview=False, force=False, renderer="auto"):
    outputs = output_paths(pptx)
    requested = [outputs[2]] + ([outputs[0]] if pdf else []) + ([outputs[1]] if preview else [])
    preflight(requested, force)
    # An unsuccessful rerender must not leave an old PDF looking like the new result.
    if force:
        for path in [outputs[2]] + ([outputs[0]] if pdf else []):
            if path.is_file():
                path.unlink()
        if preview and outputs[1].is_dir():
            for path in outputs[1].iterdir():
                if path.is_file() and re.fullmatch(r"slide-\d+\.png", path.name):
                    path.unlink()
    with tempfile.TemporaryDirectory(prefix="pchk-render-") as stage:
        if renderer == "portable" or (renderer == "auto" and os.name != "nt"):
            return render_portable(pptx, stage, outputs, pdf, preview)
        if os.name != "nt":
            raise SheetError("Рендерер PowerPoint требует Windows. Используйте --renderer portable.")
        command = ["powershell.exe", "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", str(ROOT / "scripts" / "export_powerpoint.ps1"), "-InputPath", str(pptx), "-OutputDirectory", stage]
        if pdf:
            command.append("-Pdf")
        if preview:
            command.append("-Preview")
        try:
            result = subprocess.run(command, capture_output=True, timeout=120)
        except subprocess.TimeoutExpired as exc:
            raise SheetError("Экспорт PowerPoint не завершился за 120 секунд. PPTX сохранён; проверьте зависший диалог PowerPoint.") from exc
        report_path = Path(stage) / "report.json"
        if report_path.is_file():
            report = json.loads(report_path.read_text(encoding="utf-8-sig"))
            write_json(outputs[2], report)
        else:
            report = None
        if result.returncode:
            if report and report.get("overflow"):
                names = "; ".join(f"слайд {v['slide']}: {v['shape']}" for v in report["overflow"])
                raise SheetError("PowerPoint обнаружил переполнение: " + names + ". PPTX сохранён для исправления; PDF не опубликован.")
            if renderer == "auto":
                print("PowerPoint недоступен; использую переносимый PDF-экспорт.")
                return render_portable(pptx, stage, outputs, pdf, preview)
            detail = (result.stderr or result.stdout).decode("utf-8", errors="replace").strip()
            raise SheetError("PowerPoint не выполнил проверку/экспорт. PPTX сохранён. " + detail[-1500:])
        if not report or report.get("status") != "ok":
            raise SheetError("PowerPoint не подтвердил успешную проверку")
        if pdf:
            source = Path(stage) / "sheet.pdf"
            if not source.is_file() or source.stat().st_size == 0:
                raise SheetError("PowerPoint не создал PDF")
            shutil.copy2(source, outputs[0])
            print("PDF:", outputs[0])
        if preview:
            images = sorted(Path(stage).glob("slide-*.png"))
            if len(images) != report["slides"]:
                raise SheetError("Не все слайды отрисованы")
            outputs[1].mkdir(exist_ok=True)
            for image in images:
                shutil.copy2(image, outputs[1] / image.name)
            print("PNG:", outputs[1])
        print("Проверка PowerPoint: текст помещается. Визуально проверьте PNG.")


def render_portable(pptx, stage, outputs, pdf, preview):
    from portable_pdf import export
    report = export(pptx, stage, preview)
    write_json(outputs[2], report)
    if report["overflow"]:
        names = "; ".join(f"слайд {x['slide']}: {x['shape']}" for x in report["overflow"])
        raise SheetError("Переносимый рендерер обнаружил переполнение: " + names + ". PDF не опубликован.")
    if pdf:
        shutil.copy2(Path(stage) / "sheet.pdf", outputs[0])
        print("PDF:", outputs[0])
    if preview:
        outputs[1].mkdir(exist_ok=True)
        for i in range(report["slides"]):
            name = f"slide-{i + 1}.png"
            shutil.copy2(Path(stage) / name, outputs[1] / name)
        print("PNG:", outputs[1])
    print("Переносимый экспорт: текст помещается. Проверьте PNG; шрифты PT Serif и Anton.")


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest="command", required=True)
    check = sub.add_parser("validate", help="Validate a character draft")
    check.add_argument("draft", type=Path)
    build = sub.add_parser("build", help="Fill the PChK template")
    build.add_argument("draft", type=Path)
    build.add_argument("--output-dir", type=Path)
    export = sub.add_parser("render", help="Check/export an existing PPTX using PowerPoint")
    export.add_argument("pptx", type=Path)
    for command in [build, export]:
        command.add_argument("--pdf", action="store_true")
        command.add_argument("--preview", action="store_true")
        command.add_argument("--force", action="store_true")
        command.add_argument("--renderer", choices=["auto", "powerpoint", "portable"], default="auto")
    args = parser.parse_args(argv)
    try:
        if args.command == "render":
            pptx = args.pptx.resolve()
            if not pptx.is_file() or pptx.suffix.lower() != ".pptx":
                raise SheetError("Для render нужен существующий PPTX")
            render(pptx, args.pdf, args.preview, args.force, args.renderer)
            return 0
        data = load_draft(args.draft)
        if args.command == "validate":
            print("Структура черновика корректна:", data["name"])
            return 0
        stem = re.sub(r"(?:\.draft)?\.json$", "", args.draft.name, flags=re.I)
        directory = (args.output_dir or args.draft.resolve().parent).resolve()
        pptx = directory / (stem + ".pchk.pptx")
        requested = [pptx]
        if args.pdf or args.preview:
            pdf_path, preview_path, report_path = output_paths(pptx)
            requested += [report_path] + ([pdf_path] if args.pdf else []) + ([preview_path] if args.preview else [])
        preflight(requested, args.force)
        presentation = fill(args.draft, data)
        directory.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(prefix="pchk-build-", dir=directory) as stage:
            temporary = Path(stage) / "sheet.pptx"
            presentation.save(str(temporary))
            os.replace(temporary, pptx)
        print("PPTX:", pptx)
        if args.pdf or args.preview:
            render(pptx, args.pdf, args.preview, args.force, args.renderer)
        else:
            print("Создан только PPTX; проверка переносов в PowerPoint не выполнялась.")
        return 0
    except (SheetError, OSError) as exc:
        print("ОШИБКА:", exc, file=sys.stderr)
        return 2


if __name__ == "__main__":
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
        sys.stderr.reconfigure(encoding="utf-8")
    raise SystemExit(main())
