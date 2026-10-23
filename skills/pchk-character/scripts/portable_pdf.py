"""Render the bundled character-sheet layout without an Office installation.

This is a renderer for this skill's simple text/image template, not arbitrary PowerPoint.
"""
from __future__ import annotations

import io
from pathlib import Path
from xml.sax.saxutils import escape

from PIL import Image
from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from reportlab.lib.colors import Color, HexColor
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.utils import ImageReader
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen.canvas import Canvas
from reportlab.platypus import Paragraph


def register_fonts():
    fonts = Path(__file__).resolve().parents[1] / "assets" / "fonts"
    for suffix, file in [("", "Regular"), ("-Bold", "Bold"), ("-Italic", "Italic"), ("-BoldItalic", "BoldItalic")]:
        name = "PCHK-Serif" + suffix
        if name not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont(name, str(fonts / "ptserif" / f"PT_Serif-Web-{file}.ttf")))
    pdfmetrics.registerFontFamily("PCHK-Serif", normal="PCHK-Serif", bold="PCHK-Serif-Bold", italic="PCHK-Serif-Italic", boldItalic="PCHK-Serif-BoldItalic")
    if "PCHK-Number" not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont("PCHK-Number", str(fonts / "anton" / "Anton-Regular.ttf")))


def picture(canvas, shape, page_height):
    image = Image.open(io.BytesIO(shape.image.blob))
    left, top, width, height = (v.pt for v in (shape.left, shape.top, shape.width, shape.height))
    # Use the PPTX crop metadata as a PDF clipping window; no raster image is edited.
    visible_w = 1 - shape.crop_left - shape.crop_right
    visible_h = 1 - shape.crop_top - shape.crop_bottom
    full_w, full_h = width / visible_w, height / visible_h
    canvas.saveState()
    clip = canvas.beginPath()
    clip.rect(left, page_height - top - height, width, height)
    canvas.clipPath(clip, stroke=0, fill=0)
    canvas.drawImage(ImageReader(image), left - full_w * shape.crop_left,
                     page_height - top - height - full_h * shape.crop_bottom,
                     full_w, full_h, mask="auto")
    canvas.restoreState()
    image.close()


def font_name(run):
    if run.font.name == "Impact":
        return "PCHK-Number"
    bold, italic = bool(run.font.bold), bool(run.font.italic)
    return "PCHK-Serif" + ("-BoldItalic" if bold and italic else "-Bold" if bold else "-Italic" if italic else "")


def text(canvas, shape, page_height, page_number):
    tf = shape.text_frame
    x = shape.left.pt + tf.margin_left.pt
    top = shape.top.pt + tf.margin_top.pt
    width = shape.width.pt - tf.margin_left.pt - tf.margin_right.pt
    bottom = shape.top.pt + shape.height.pt - tf.margin_bottom.pt
    overflow = []
    for paragraph in tf.paragraphs:
        runs = list(paragraph.runs)
        if not paragraph.text:
            # PowerPoint's empty paragraphs are full blank lines, e.g. in a backstory.
            top += 16 * 1.22 * 1.03
            continue
        size = runs[0].font.size.pt if runs[0].font.size else 9.4
        if runs[0].font.name == "Impact":
            value = paragraph.text
            face = font_name(runs[0])
            span = pdfmetrics.stringWidth(value, face, size)
            if span > width:
                overflow.append(dict(slide=page_number, shape=shape.name, reason="number width"))
            tx = x + (width - span) / 2
            canvas.saveState()
            canvas.setLineWidth(0.4)
            canvas.setStrokeColor(HexColor("#242222"))
            canvas.setFillColorRGB(1, 1, 1)
            obj = canvas.beginText(tx, page_height - top - size * 0.95)
            obj.setFont(face, size)
            obj.setTextRenderMode(2)
            obj.textOut(value)
            canvas.drawText(obj)
            canvas.restoreState()
            top += size * 1.2
            continue
        fragments = []
        for run in runs:
            value = escape(run.text).replace("\n", "<br/>")
            fragments.append(f'<font name="{font_name(run)}" size="{run.font.size.pt if run.font.size else size}">{value}</font>')
        spacing = paragraph.line_spacing
        leading = spacing.pt if hasattr(spacing, "pt") else size * 1.22 * (spacing or 1.0)
        style = ParagraphStyle("sheet", fontName="PCHK-Serif", fontSize=size, leading=leading,
                               textColor=Color(36 / 255, 34 / 255, 34 / 255),
                               alignment={PP_ALIGN.CENTER: TA_CENTER, PP_ALIGN.RIGHT: TA_RIGHT}.get(paragraph.alignment, TA_LEFT),
                               splitLongWords=True, allowWidows=0, allowOrphans=0)
        block = Paragraph("".join(fragments), style)
        _, height = block.wrap(width, page_height * 100)
        top += paragraph.space_before.pt if paragraph.space_before else 0
        if top + height > bottom + 1.2:
            overflow.append(dict(slide=page_number, shape=shape.name, actual_bottom=round(top + height, 2), allowed_bottom=round(bottom, 2)))
        block.drawOn(canvas, x, page_height - top - height)
        top += height + (paragraph.space_after.pt if paragraph.space_after else 0)
    return overflow


def export(pptx, directory, preview=False):
    register_fonts()
    directory = Path(directory)
    presentation = Presentation(str(pptx))
    width, height = presentation.slide_width.pt, presentation.slide_height.pt
    pdf_path = directory / "sheet.pdf"
    canvas = Canvas(str(pdf_path), pagesize=(width, height), pageCompression=1)
    canvas.setTitle(presentation.core_properties.title)
    overflow = []
    for page_number, slide in enumerate(presentation.slides, 1):
        backgrounds = [s for s in list(slide.slide_layout.slide_master.shapes) + list(slide.slide_layout.shapes) if hasattr(s, "image")]
        for shape in backgrounds + list(slide.shapes):
            if hasattr(shape, "image"):
                picture(canvas, shape, height)
            elif shape.has_text_frame and shape.text:
                overflow.extend(text(canvas, shape, height, page_number))
        canvas.showPage()
    canvas.save()
    # Report a field once even if several paragraphs exceed its available space.
    overflow = list({(item["slide"], item["shape"]): item for item in overflow}.values())
    report = dict(status="overflow" if overflow else "ok", slides=len(presentation.slides),
                  renderer="ReportLab (PT Serif / Anton)", overflow=overflow,
                  scope="Bundled PChK layout only; verify PNG visually. Fonts substitute Cambria/Impact.")
    if preview and not overflow:
        import pypdfium2 as pdfium
        doc = pdfium.PdfDocument(str(pdf_path))
        try:
            for i in range(len(doc)):
                page = doc[i]
                bitmap = page.render(scale=1800 / width)
                bitmap.to_pil().save(directory / f"slide-{i + 1}.png")
                bitmap.close()
                page.close()
        finally:
            doc.close()
    return report
