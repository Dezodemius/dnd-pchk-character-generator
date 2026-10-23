import copy
import hashlib
import importlib.util
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
import zipfile

from pptx import Presentation
import pypdfium2 as pdfium

ROOT = Path(__file__).resolve().parents[1]
SKILL = ROOT / "skills" / "pchk-character"
sys.path.insert(0, str(SKILL / "scripts"))
import pchk


class SkillTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.work = Path(self.temporary.name)
        self.data = json.loads((SKILL / "assets/example.draft.json").read_text(encoding="utf-8"))

    def draft(self, data=None, name="hero.draft.json"):
        path = self.work / name
        pchk.write_json(path, data if data is not None else self.data)
        return path

    def test_validation_rejects_scores_booleans_and_unknown_fields(self):
        for key, value in [("hit_points", True), ("ability_modifiers", dict(self.data["ability_modifiers"], str=16)), ("surprise", "typo"), ("portrait_focus", [float("nan"), 0])]:
            data = copy.deepcopy(self.data)
            data[key] = value
            with self.subTest(key=key), self.assertRaises(pchk.SheetError):
                pchk.load_draft(self.draft(data))

    def test_portable_end_to_end(self):
        draft = self.draft()
        code = pchk.main(["build", str(draft), "--pdf", "--preview", "--renderer", "portable"])
        self.assertEqual(code, 0)
        pptx = self.work / "hero.pchk.pptx"
        deck = Presentation(pptx)
        self.assertEqual(len(deck.slides), 2)
        self.assertEqual(pchk.shape_by_name(deck.slides[0], "Сила").text, "+3")
        self.assertEqual(pchk.shape_by_name(deck.slides[0], "Харизма").text, "-1")
        self.assertEqual(pchk.shape_by_name(deck.slides[0], "Класс брони").text, "16/18")
        document = pdfium.PdfDocument(str(pptx.with_suffix(".pdf")))
        try:
            self.assertEqual(len(document), 2)
            page = document[0]
            text = page.get_textpage()
            self.assertIn("Марта Медный Ключ", text.get_text_range())
            text.close()
            page.close()
        finally:
            document.close()
        self.assertEqual(len(list(pptx.with_suffix(".preview").glob("*.png"))), 2)
        original = hashlib.sha256(pptx.read_bytes()).hexdigest()
        self.assertEqual(pchk.main(["build", str(draft)]), 2)
        self.assertEqual(hashlib.sha256(pptx.read_bytes()).hexdigest(), original)

    def test_overflow_does_not_publish_stale_pdf(self):
        draft = self.draft()
        args = ["build", str(draft), "--pdf", "--renderer", "portable", "--force"]
        self.assertEqual(pchk.main(args), 0)
        data = copy.deepcopy(self.data)
        data["backstory"] *= 30
        self.draft(data)
        self.assertEqual(pchk.main(args), 2)
        self.assertFalse((self.work / "hero.pchk.pdf").exists())
        report = json.loads((self.work / "hero.pchk.render-report.json").read_text(encoding="utf-8"))
        self.assertEqual(report["status"], "overflow")

    def test_special_block_and_relative_portrait(self):
        # Use a real bundled image; do not require a network-generated portrait in CI.
        with zipfile.ZipFile(SKILL / "assets/character-sheet.pptx") as archive:
            image = next(x for x in archive.namelist() if x.startswith("ppt/media/") and x.endswith(".png"))
            (self.work / "portrait.png").write_bytes(archive.read(image))
        self.data["portrait"] = "portrait.png"
        self.data["special"] = dict(title="Ключ", charges=3, recovery="На рассвете", effects=[dict(name="Память", uses="1 заряд", text="Показывает исчезнувшую дверь.")])
        draft = self.draft()
        deck = pchk.fill(draft, pchk.load_draft(draft))
        self.assertIn("заряды: 3", pchk.shape_by_name(deck.slides[0], "Заголовок особого ресурса").text)
        portrait = next(x for x in deck.slides[1].shapes if x.name.startswith("Портрет"))
        self.assertGreater(portrait.crop_left + portrait.crop_right + portrait.crop_top + portrait.crop_bottom, 0)


if __name__ == "__main__":
    unittest.main()
